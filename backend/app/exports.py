"""Screening result exports — CSV, Excel (openpyxl), PDF report (ReportLab)."""

import csv
import io
from datetime import datetime

EXPORT_COLUMNS = [
    ("rank", "Rank"),
    ("name", "Candidate"),
    ("email", "Email"),
    ("phone", "Phone"),
    ("score_pct", "Match Score %"),
    ("band", "Verdict"),
    ("experience_years", "Experience (yrs)"),
    ("matched_skills", "Matched Skills"),
    ("missing_skills", "Missing Skills"),
    ("status", "Status"),
    ("filename", "Resume File"),
]


def _rows(session, candidates):
    out = []
    for rank, c in enumerate(candidates, start=1):
        out.append({
            "rank": rank,
            "name": c.name or "(name not found)",
            "email": c.email or "",
            "phone": c.phone or "",
            "score_pct": round(c.score * 100, 1),
            "band": c.band,
            "experience_years": c.experience_years,
            "matched_skills": "; ".join(c.matched_skills or []),
            "missing_skills": "; ".join(c.missing_skills or []),
            "status": c.status,
            "filename": c.filename,
        })
    return out


def build_csv(session, candidates) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([f"Screening report — {session.job_title} @ {session.job_company}".strip(" @")])
    w.writerow([f"Generated {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')} · "
                f"{len(candidates)} candidates screened"])
    w.writerow([])
    w.writerow([label for _, label in EXPORT_COLUMNS])
    for r in _rows(session, candidates):
        w.writerow([r[key] for key, _ in EXPORT_COLUMNS])
    return buf.getvalue()


def build_xlsx(session, candidates) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Ranking"
    header_fill = PatternFill("solid", fgColor="1F2A44")
    header_font = Font(bold=True, color="FFFFFF")
    title = f"{session.job_title} @ {session.job_company}".strip(" @")
    ws.append([f"Screening report — {title}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Generated {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')} · "
               f"{len(candidates)} candidates screened"])
    ws.append([])
    ws.append([label for _, label in EXPORT_COLUMNS])
    for cell in ws[4]:
        cell.fill = header_fill
        cell.font = header_font
    for r in _rows(session, candidates):
        ws.append([r[key] for key, _ in EXPORT_COLUMNS])
    for col in ws.columns:
        width = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(width + 2, 45)
    for row in ws.iter_rows(min_row=5):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    # Job sheet
    job = session.job_snapshot or {}
    wj = wb.create_sheet("Job")
    wj.append(["Field", "Value"])
    for cell in wj[1]:
        cell.fill = header_fill
        cell.font = header_font
    wj.append(["Title", session.job_title])
    wj.append(["Company", session.job_company])
    wj.append(["Required skills", ", ".join(job.get("required_skills", []))])
    wj.append(["Preferred skills", ", ".join(job.get("preferred_skills", []))])
    wj.append(["Min experience (yrs)", job.get("min_experience_years", 0)])
    wj.append(["Education requirements", job.get("education_requirements", "")])
    wj.append(["Job description", job.get("jd_text", "")])
    wj.column_dimensions["A"].width = 24
    wj.column_dimensions["B"].width = 80

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf(session, candidates) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=16 * mm, bottomMargin=16 * mm)
    styles = getSampleStyleSheet()
    h1, h2, body = styles["Heading1"], styles["Heading2"], styles["BodyText"]
    small = ParagraphStyle("Small", parent=styles["Normal"], fontSize=8, leading=11,
                           textColor=colors.HexColor("#555555"))
    story = []
    title = f"{session.job_title} @ {session.job_company}".strip(" @")
    story.append(Paragraph("Resume Screening Report", h1))
    story.append(Paragraph(title, h2))
    story.append(Paragraph(
        f"Generated {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')} · "
        f"{len(candidates)} candidate(s) screened · "
        f"Session #{session.id}", small))
    story.append(Spacer(1, 6 * mm))
    story.append(Paragraph(
        "Scores are similarity assessments to assist hiring decisions — never hiring "
        "decisions. Names, colleges and contact details are never used in scoring, "
        "and no candidate is ever rejected automatically.", small))
    story.append(Spacer(1, 6 * mm))

    story.append(Paragraph("Ranking", h2))
    tdata = [["#", "Candidate", "Score", "Verdict", "Exp (yrs)", "Status"]]
    for r in _rows(session, candidates):
        tdata.append([str(r["rank"]), r["name"], f"{r['score_pct']}%",
                      r["band"], str(r["experience_years"]), r["status"]])
    t = Table(tdata, colWidths=[12 * mm, 52 * mm, 22 * mm, 32 * mm, 22 * mm, 32 * mm],
              repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2A44")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F5FA")]),
    ]))
    story.append(t)
    story.append(Spacer(1, 8 * mm))

    story.append(Paragraph("Candidate skill evidence", h2))
    for r in _rows(session, candidates):
        story.append(Paragraph(f"<b>#{r['rank']} {r['name']}</b> — {r['score_pct']}% ({r['band']})", body))
        story.append(Paragraph(f"Matched: {r['matched_skills'] or '—'}", small))
        story.append(Paragraph(f"Missing: {r['missing_skills'] or '—'}", small))
        story.append(Spacer(1, 3 * mm))

    doc.build(story)
    return buf.getvalue()
